"""Generate data/template_records.xlsx - the school's fill-in dataset template.

Matches the exact 9 original attributes of records.xlsx so the upload/convert
flow (ml/convert_history.py) can read it directly. Includes an Attribute Guide
sheet that tells panelists which attributes are predictors and which is the
target (Prediction).

Run:  python make_records_template.py
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from config import Config

HEADERS = [
    "Name", "Province", "GWA", "Units", "Failed Subject",
    "Annual Family Income", "Attendance", "Socio Economic Status", "Prediction",
]

EXAMPLES = [
    ["Juan Dela Cruz", "Zambales", 1.25, 18, 0, 150000, 97, "Low", "Retained"],
    ["Maria Santos", "Zambales", 2.10, 20, 1, 185000, 89, "Lower-Middle", "At-Risk"],
    ["Carlo Bautista", "Bataan", 1.80, 21, 0, 120000, 95, "Low", "Retained"],
]

GUIDE = [
    ("1", "Name", "Student's full name", "Identifier (not used as a predictor)", "Juan Dela Cruz"),
    ("2", "Province", "Student's province", "Informational (not used as a predictor)", "Zambales"),
    ("3", "GWA", "General weighted average (lower is better)", "Predictor (feature)", "1.25"),
    ("4", "Units", "Units enrolled for the term", "Predictor (feature)", "18"),
    ("5", "Failed Subject", "Number of failed subjects", "Predictor (feature)", "0"),
    ("6", "Annual Family Income", "Family's annual income (PHP)", "Predictor (feature)", "150000"),
    ("7", "Attendance", "Class attendance rate (%)", "Predictor (feature)", "95"),
    ("8", "Socio Economic Status", "Low / Lower-Middle / Middle / Upper-Middle", "Predictor (feature)", "Low"),
    ("9", "Prediction", "Retention outcome of the scholar", "Target / Output", "Retained / At-Risk"),
]

HEADER_FILL = PatternFill("solid", fgColor="1B4F8A")
HEADER_FONT = Font(bold=True, color="FFFFFF")
GUIDE_FILL = PatternFill("solid", fgColor="E3F0FF")


def main():
    out = os.path.join(Config.DATA_DIR, "template_records.xlsx")
    wb = Workbook()

    ws = wb.active
    ws.title = "Records"
    ws.append(HEADERS)
    for i, cell in enumerate(ws[1], 1):
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for row in EXAMPLES:
        ws.append(row)
    ws.freeze_panes = "A2"

    widths = [22, 14, 10, 10, 14, 20, 12, 22, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    pred_dv = DataValidation(type="list", formula1='"Retained,At-Risk"', allow_blank=True)
    ws.add_data_validation(pred_dv)
    pred_dv.add(f"I2:I1000")
    socio_dv = DataValidation(
        type="list",
        formula1='"Low,Lower-Middle,Middle,Upper-Middle"',
        allow_blank=True,
    )
    ws.add_data_validation(socio_dv)
    socio_dv.add(f"H2:H1000")

    guide = wb.create_sheet("Attribute Guide")
    guide.append(["#", "Attribute", "Meaning", "Role in Model", "Example"])
    for i, cell in enumerate(guide[1], 1):
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row in GUIDE:
        guide.append(row)
    guide.append([])
    guide.append(["Note",
                  "The processed model dataset (scholar_data.csv) may contain additional "
                  "standardized fields (e.g. scholarship_type, semester_performance, "
                  "academic_year) introduced during preprocessing. These are NOT direct "
                  "columns of this file and should not be interpreted as separately "
                  "collected data."])
    guide["A12"].font = Font(bold=True)
    guide["B12"].alignment = Alignment(wrap_text=True)
    guide.row_dimensions[12].height = 45

    for i, w in enumerate([4, 26, 42, 42, 22], 1):
        guide.column_dimensions[get_column_letter(i)].width = w
    for row in guide.iter_rows(min_row=2, max_row=11):
        for cell in row:
            cell.alignment = Alignment(vertical="top")

    wb.save(out)
    print(f"Template written -> {out}")


if __name__ == "__main__":
    main()