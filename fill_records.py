"""Clean up data/records.xlsx: remove fully-blank rows and fix erroneous values.

Findings in the school's file:
  - 7,690 raw rows, of which 18 are completely empty (no data in any column)
  - Province contains '5' in 20 rows (erroneous) instead of 'ZAMBALES'
  - After cleaning: exactly 7,672 complete records (matches the training dataset)

The original file is backed up to data/records_backup.xlsx first.

Run:  python fill_records.py
"""
import os
import shutil
import sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from config import Config

PROVINCE_VALUE = "ZAMBALES"
BAD_PROVINCE = ("5", "5.0", "n/a", "N/A", "NA", "")


def _is_blank(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def main():
    src = os.path.join(Config.DATA_DIR, "records.xlsx")
    if not os.path.exists(src):
        raise SystemExit(f"Missing source file: {src}")
    backup = os.path.join(Config.DATA_DIR, "records_backup.xlsx")
    shutil.copy2(src, backup)
    print(f"Backup saved -> {backup}")

    wb = load_workbook(src)
    ws = wb.active
    header = [str(c.value) if c.value is not None else "" for c in ws[1]]
    ncol = len(header)

    removed = 0
    fixed = 0
    prov_col = header.index("Province") + 1

    # Delete fully-blank rows (iterate bottom-up so indices stay valid after delete).
    for r in range(ws.max_row, 1, -1):
        cells = [ws.cell(row=r, column=c).value for c in range(1, ncol + 1)]
        if all(_is_blank(v) for v in cells):
            ws.delete_rows(r, 1)
            removed += 1

    # Normalize erroneous Province values.
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=prov_col).value
        if v is not None and str(v).strip() in BAD_PROVINCE:
            ws.cell(row=r, column=prov_col).value = PROVINCE_VALUE
            fixed += 1

    wb.save(src)
    print(f"Cleaned {src}: removed {removed} blank row(s), fixed {fixed} Province value(s).")
    print(f"Data rows now: {ws.max_row - 1}")
    print(f"Province distinct values: "
          f"{sorted({ws.cell(row=r, column=prov_col).value for r in range(2, ws.max_row + 1)})}")


if __name__ == "__main__":
    main()